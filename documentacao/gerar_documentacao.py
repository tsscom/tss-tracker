from pathlib import Path
from html import escape
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import letter

OUT = Path(__file__).resolve().parent
DATE = '30 setembro 2026'
COMPANY = 'TSS Transportes Sergiu Sandu, Unipessoal Lda.'
WIDTH = 468
PAGES = {}

def p(text): return ('p', text)
def h(text): return ('h', text)
def b(*texts): return ('b', list(texts))
def table(headers, rows, widths): return ('table', headers, rows, widths)
def page(title, *items): return {'title':title, 'items':list(items)}

roles = [
 ('gerente','Sergiu Sandu','Aprova cotações, crédito, documentos internos, cancelamentos, orçamento e libertação de oficina. Coordena a validação e assume as exceções.'),
 ('administrador','Por atribuir','Cria e desativa identidades, atribui perfis, mantém configurações e cópias. Este perfil, por si só, não aprova decisões comerciais.'),
 ('comercial','Por atribuir','Regista clientes, oportunidades e cotações. Documenta a aceitação do cliente e solicita alterações de crédito.'),
 ('operacoes','Por atribuir','Converte a cotação aceite em serviço. Planeia, atribui motorista e veículo, despacha e acompanha a entrega.'),
 ('motorista','Por atribuir','Consulta apenas os serviços que lhe são atribuídos. Regista execução, incidentes e prova de entrega.'),
 ('parque','Por atribuir','Controla entrada, posição, movimentação e saída de veículos no parque.'),
 ('oficina','Por atribuir','Abre ordens de trabalho, estima custos, executa trabalhos aprovados e regista conclusão e evidência técnica.'),
 ('financeiro','Por atribuir','Prepara documentos internos e pedidos de aprovação. Regista recebimentos, reconcilia saldos e acompanha cobranças.'),
 ('auditor','Por atribuir','Consulta evidência e histórico autorizado. Não cria, altera, aprova nem elimina registos operacionais.')
]

requirements = [
 ('REQ01','Identidade','Cada pessoa tem conta própria; a configuração inicial cria apenas a conta de Sergiu.','Local','US01'),
 ('REQ02','Autenticação','Autenticação verificada no servidor, palavras passe derivadas com sal, saída e expiração de sessão.','Local','US02'),
 ('REQ03','Autorização','Validar perfil e alcance de cada operação no servidor; motorista limitado aos serviços atribuídos.','Local','US03 US04'),
 ('REQ04','Aprovações','Gerente aprova; autoaprovação normal rejeitada; exceção do proprietário exige motivo com 20 caracteres e histórico.','Local','US05'),
 ('REQ05','Histórico','Guardar ator, ação, entidade, instante e motivo das ações e decisões; exportação autorizada.','Local','US06'),
 ('REQ06','Segurança em produção','Implementar HTTPS, MFA para perfis privilegiados, recuperação segura, monitorização e gestão de segredos.','Produção','US38 US39'),
 ('REQ07','Clientes','Ficha única com identificação, contactos, moradas, estado e observações; rejeitar duplicações evidentes.','Local parcial','US07'),
 ('REQ08','Relação comercial','Oportunidades com responsável, etapa, tarefa seguinte e histórico de contactos ligados ao cliente.','Produção','US08'),
 ('REQ09','Crédito','Alterar limite por pedido e decisão do gerente; valores e motivos registados sem limites de aprovação inventados.','Local','US09'),
 ('REQ10','Cotações','Cliente, recolha, entrega, datas e preço em euros; enviar a aprovação antes da aceitação e conversão.','Local','US10 US11'),
 ('REQ11','Aceitação','Aceitação com referência do cliente; criar um serviço a partir da cotação aceite sem duplicar a conversão.','Local','US12'),
 ('REQ12','Serviço','Referência única e ligação ao cliente; datas, recolha, entrega, motorista, veículo, valor, estado e pagamento.','Local','US13'),
 ('REQ13','Planeamento','Operações atribui recursos válidos; verificar conflitos de período e indisponibilidade da oficina.','Local parcial','US14'),
 ('REQ14','Despacho','Despachar apenas serviço agendado com recursos atribuídos e veículo disponível; manter hora do despacho.','Local','US15'),
 ('REQ15','Execução','Transições controladas de Agendado para Em curso e Concluído; registar prova de entrega antes de concluir.','Local','US16 US17'),
 ('REQ16','Incidentes','Registar avaria, atraso, dano e falha de entrega com responsável, evidência e plano de resolução.','Produção','US18'),
 ('REQ17','Cancelamento','Pedir cancelamento com motivo; somente decisão do gerente altera o estado final.','Local','US19'),
 ('REQ18','Pesquisa','Pesquisar e combinar filtros; exportar apenas os registos autorizados e correspondentes aos filtros.','Local','US20'),
 ('REQ19','Parque','Registar entrada, veículo, portão, instante e posição livre; uma presença ativa por veículo.','Local parcial','US21'),
 ('REQ20','Movimentação','Transferir presença entre posições livres; guardar origem, destino, ator e instante.','Local','US22'),
 ('REQ21','Saída','Registar saída uma vez; não admitir saída operacional de veículo bloqueado pela oficina.','Local','US23'),
 ('REQ22','Capacidade do parque','Configurar posições, reservas e capacidade; impedir ocupações simultâneas em posição exclusiva.','Produção','US24'),
 ('REQ23','Oficina','Ordem ligada ao veículo com problema, estimativa, prioridade e estado; bloquear despacho enquanto ativa.','Local parcial','US25'),
 ('REQ24','Trabalhos','Submeter orçamento ao gerente; executar apenas ordem aprovada e registar conclusão e custo efetivo.','Local','US26 US27'),
 ('REQ25','Libertação','Gerente liberta veículo após conclusão técnica e decisão documentada; separar de quem concluiu salvo exceção do proprietário.','Local','US28'),
 ('REQ26','Manutenção e peças','Planos por tempo ou quilometragem, stock, consumos e alertas com histórico por veículo.','Produção','US29'),
 ('REQ27','Documento interno','Criar uma minuta por serviço concluído com prova de entrega; valor ligado ao preço do serviço.','Local','US30'),
 ('REQ28','Aprovação financeira','Submeter e aprovar antes de marcar emissão interna; correções preservam rastreabilidade.','Local','US31'),
 ('REQ29','Recebimentos','Valor positivo, referência única e montante não superior ao saldo; suportar recebimento parcial e integral.','Local','US32'),
 ('REQ30','Cobranças','Mostrar saldo e vencimento; lista de valores vencidos e seguimento de contactos com o cliente.','Local e produção','US33'),
 ('REQ31','Faturação externa','Ligar a faturação real e contabilidade através de integração validada com o contabilista; piloto não emite documento fiscal.','Produção','US34'),
 ('REQ32','Persistência e migração','Preservar serviços existentes e cópia anterior à migração; marcar informação histórica que exige revisão.','Local','US35'),
 ('REQ33','Cópias e recuperação','Cópia exportável restrita e procedimento de restauro ensaiado, incluindo contas, histórico e relações.','Local e produção','US36'),
 ('REQ34','Qualidade','Teste do ciclo completo, negativas por perfil, dupla decisão, duplicação de pagamento e conflito de recursos.','Local e produção','US37'),
 ('REQ35','Proteção de dados','Definir finalidade, acesso, conservação, pedidos de titulares e eliminação controlada antes de dados reais em escala.','Produção','US40'),
 ('REQ36','Concorrência','Migrar para base transacional, controlo de versão e unicidade; nenhuma perda sob alterações simultâneas.','Produção','US41'),
 ('REQ37','Integrações','Conectores de telemática, documentos, correio e banca com reprocessamento seguro e reconciliação.','Produção','US42'),
 ('REQ38','Aceitação operacional','Treinar cada perfil, validar procedimentos e executar piloto acompanhado antes de decisão de entrada em produção.','Produção','US43'),
 ('REQ39','Indicadores','Painéis por período para pontualidade, saldo, ocupação e custo de oficina; definições e fonte acessíveis.','Produção','US44')
]

PAGES['01_Requisitos_e_governacao'] = [
 page('Requisitos e governação da plataforma TSS',
   p('Este documento define o sistema integrado de clientes, transportes, parque, oficina e recebimentos da TSS. Serve de referência para o desenvolvimento e para Sergiu Sandu validar a forma como cada decisão é registada. A versão local permite ensaiar os fluxos com dados de demonstração; a utilização empresarial em produção exige os trabalhos identificados neste documento.'),
   table(['Controlo','Valor'],[['Entidade',COMPANY],['Localidade','Alenquer Portugal'],['Responsável de aprovação','Sergiu Sandu'],['Versão','1.0'],['Data',DATE],['Utilização','Validação local de fluxos e preparação do produto empresarial']],[140,328]),
   h('Âmbito funcional'),
   b('CRM  gestão de clientes, crédito, cotações e aceitação comercial.','TMS  planeamento, afetação de recursos, despacho, execução e prova de entrega.','YMS  entrada, posição, movimentos e saída do parque.','Oficina  orçamento, aprovação, execução, conclusão e libertação de veículos.','Financeiro  minutas internas, aprovação, recebimentos e saldos ligados ao serviço.'),
   h('Decisões para o piloto'),
   p('A única pessoa identificada é Sergiu Sandu. Os restantes perfis representam funções a atribuir, sem nomes, contactos ou credenciais predefinidos. Todas as decisões comerciais e financeiras sujeitas a aprovação pertencem ao gerente; não são criados escalões monetários de aprovação. O administrador gere o sistema e as identidades.'),
   p('A regra normal exige que o decisor seja uma pessoa diferente do requerente. Como o ensaio local pode ser operado por Sergiu sozinho, existe a opção explícita de exceção do proprietário. O sistema exige motivo com pelo menos 20 caracteres e grava ator, requerente, data e justificação. Em produção com equipa, deve aplicar-se separação de funções por omissão.')
 ),
 page('Responsabilidades e acesso',
   p('Um perfil determina o que uma pessoa pode fazer. No piloto, cada conta tem um perfil. O gerente possui as funções de negócio necessárias ao ensaio individual. A tabela descreve responsabilidades; as atribuições efetivas devem ser registadas na lista de utilizadores e revistas por Sergiu.'),
   table(['Perfil','Titular','Responsabilidade'],roles,[80,85,303]),
   h('Princípios de atribuição'),
   b('Criar contas nominais apenas para pessoas autorizadas. Não partilhar a conta de Sergiu.','Atribuir o perfil necessário à tarefa aprovada. O perfil administrador não concede aprovação de negócio.','Associar a conta motorista ao serviço para limitar os registos visíveis.','Desativar o acesso quando termina a responsabilidade e rever sessões ativas.','Guardar pedido, aprovação, data e responsável de cada atribuição de perfil.'),
   p('A revisão de acessos deve verificar identidade, funções atuais, alcance e necessidade de exportação. Propõe-se uma revisão mensal no piloto e antes de qualquer aumento da equipa; a cadência definitiva pertence a Sergiu.')
 ),
 page('Regras de aprovação e exceções',
   p('A aprovação confirma uma versão concreta do registo. O piloto recusa alterações com versão desatualizada e bloqueia a edição comercial depois da aprovação. Na produção, acrescentar revisão formal: alterar campos relevantes deve invalidar a aprovação e originar novo pedido, conservando a versão anterior.'),
   table(['Decisão','Quem prepara','Quem decide','Evidência'],[
    ['Cotação','Comercial','Gerente','Preço, datas, locais e motivo da decisão'],
    ['Limite de crédito','Comercial com informação financeira','Gerente','Limite atual e proposto, justificação'],
    ['Cancelamento de serviço','Operações','Gerente','Motivo e tratamento de custos já incorridos'],
    ['Documento interno','Financeiro','Gerente','Serviço concluído, prova e valor'],
    ['Orçamento de oficina','Oficina','Gerente','Veículo, problema, trabalhos e estimativa'],
    ['Libertação de veículo','Oficina conclui','Gerente','Conclusão técnica, custo e decisão']
   ],[106,97,65,200]),
   h('Sequência de uma decisão'),
   b('O responsável grava o registo e submete o pedido, identificando o motivo.','O gerente verifica a evidência e aprova ou rejeita com justificação.','A decisão regista ator, requerente, entidade, versão, data e resultado.','Após rejeição, o responsável corrige e volta a submeter segundo os estados permitidos.','Pedidos duplicados ou já decididos não devem executar novamente a operação.'),
   h('Exceção do proprietário no ensaio local'),
   p('Sergiu deve selecionar a opção de exceção e explicar por que está a acumular preparação e decisão. Uma indicação genérica como aprovado não cumpre os 20 caracteres mínimos. Exemplo adequado: Estou a validar sozinho o fluxo local desta cotação. A exceção fica visível no histórico e não cria uma segunda pessoa.'),
   p('Não há dispensa implícita de prova de entrega, saldo de pagamento, bloqueio de oficina ou controlo de estados. Uma exceção de separação de funções não deve remover as restantes validações.')
 ),
 page('Processo do cliente ao recebimento',
   p('O identificador do cliente liga a cotação, o serviço, a prova de entrega, a minuta e os recebimentos. Esta cadeia permite localizar a origem de um valor e verificar a decisão que permitiu avançar.'),
   table(['Etapa','Responsável','Condição para avançar','Registo resultante'],[
    ['1 Registar cliente','Comercial','Identificação e contacto verificados','Ficha de cliente'],
    ['2 Rever crédito','Financeiro e gerente','Decisão registada quando se pede alteração','Limite aprovado'],
    ['3 Preparar cotação','Comercial','Locais, datas e valor completos','Cotação em aprovação'],
    ['4 Aprovar cotação','Gerente','Evidência revista e decisão válida','Cotação aprovada'],
    ['5 Registar aceitação','Comercial','Referência do cliente','Cotação aceite'],
    ['6 Criar e planear','Operações','Recursos válidos e disponíveis','Serviço agendado'],
    ['7 Despachar e executar','Operações e motorista','Despacho autorizado e ausência de bloqueio','Serviço em curso'],
    ['8 Confirmar entrega','Motorista e operações','Destinatário e referência da prova','Serviço concluído'],
    ['9 Preparar e aprovar minuta','Financeiro e gerente','Conclusão e prova existentes','Minuta aprovada'],
    ['10 Registar emissão interna','Financeiro','Minuta aprovada','Documento interno emitido'],
    ['11 Reconciliar recebimento','Financeiro','Referência única e montante válido','Saldo parcial ou nulo']
   ],[109,100,153,106]),
   h('Processos de suporte'),
   p('O parque controla a presença física e deve refletir os movimentos reais. A oficina torna o veículo indisponível para despacho desde a abertura de uma ordem ativa até à rejeição ou libertação. O processo comercial continua ligado ao serviço mesmo quando existe cancelamento, avaria ou atraso.'),
   p('Emitida é uma etapa interna do piloto. A minuta não é documento fiscal e não substitui a faturação utilizada pela empresa. Na produção, a confirmação da emissão real virá do sistema de faturação externo e será ligada ao serviço.')
 ),
 page('Estados e integridade dos registos',
   table(['Objeto','Fluxo principal','Condições principais'],[
    ['Cotação','Rascunho > Em aprovação > Aprovada > Aceite > Convertida','Gerente decide; aceitação tem referência; conversão cria um serviço. Rejeitada exige correção.'],
    ['Serviço','Agendado > Em curso > Concluído','Despacho fica datado; recursos válidos; conclusão tem destinatário e referência de prova. Cancelado depende de aprovação.'],
    ['Oficina','Rascunho > Em aprovação > Aprovada > Em curso > Concluída > Libertada','Ordem ativa bloqueia; gerente aprova orçamento e liberta veículo após conclusão.'],
    ['Documento interno','Rascunho > Em aprovação > Aprovada > Emitida > Paga','Uma minuta por serviço elegível; parcial mantém Emitida; saldo nulo determina Paga.'],
    ['Parque','Entrada > Presença > Movimento > Saída','Uma presença ativa por veículo; posição exclusiva não admite sobreposição.']
   ],[91,172,205]),
   h('Modelo de dados e relações'),
   b('Cliente possui cotações, serviços e limite de crédito. A identificação é estável mesmo que o nome mude.','Cotação conserva cliente, percurso, datas, valor e decisões. O serviço conserva a referência da cotação convertida.','Serviço liga motorista e veículo e possui prova de entrega. Registos antigos sem prova devem ser revistos antes de faturação interna.','Documento interno liga um serviço e vários recebimentos. O saldo resulta do valor menos a soma dos recebimentos válidos.','Veículo liga movimentos de parque e ordens de oficina. Um bloqueio ativo deve ser consultado no despacho.','Utilizador liga perfis e ações do histórico. Desativar uma conta não remove o seu histórico.'),
   h('Valores e datas'),
   p('O piloto guarda montantes em cêntimos e apresenta euros com duas casas decimais. As datas de transporte representam dias completos; o controlo de sobreposição é realizado no despacho e início. O calendário usa Europe Lisbon. Na produção, definir janelas horárias e validar sobreposições também durante o planeamento.')
 ),
 page('Segurança e continuidade',
   p('Autenticação confirma a conta; autorização verifica cada ação e o alcance do registo. O servidor deve negar operações sem permissão mesmo que alguém altere a interface ou chame diretamente a API. O alcance do motorista inclui somente serviços atribuídos. Estas regras seguem as orientações OWASP sobre autorização [2].'),
   h('Controlos da versão local'),
   b('Configuração inicial sem credenciais partilhadas de demonstração; palavra passe definida por Sergiu.','Credenciais derivadas com sal e nunca guardadas em texto simples; erros de entrada sem revelar contas.','Sessão no servidor, saída, expiração e verificação de acesso em operações protegidas.','Histórico das alterações e decisões; cópias completas tratadas como informação confidencial.','Aplicação ligada ao computador local; acesso remoto exige trabalho adicional e nova validação.'),
   h('Trabalho necessário para produção'),
   p('Antes de acesso remoto, exigir HTTPS e rever cookies, proteção contra pedidos indevidos e limitação de tentativas. Acrescentar MFA aos perfis privilegiados, recuperação segura de conta e revogação de sessões após incidentes. A orientação de autenticação OWASP fundamenta estes requisitos propostos [1].'),
   p('Rever os parâmetros da função de derivação de palavras passe na infraestrutura escolhida. OWASP recomenda algoritmos apropriados e parâmetros ajustados ao custo de ataque; a implementação do piloto deve ser medida e revista antes de produção [3].'),
   h('Cópias e recuperação'),
   p('O administrador prepara uma cópia autorizada da base completa, identifica a data e guarda-a em localização restrita. No piloto, testar a recuperação numa cópia separada. Em produção, propõe-se cópia diária, cópia independente e teste mensal de restauro. Estas frequências são propostas operacionais, sem compromisso de disponibilidade.'),
   p('Proposta a validar por Sergiu: perda máxima tolerada de dados de um dia e recuperação em quatro horas. Confirmar os objetivos depois de medir o tempo real de cópia e restauro, escolher a infraestrutura e definir o responsável de suporte.')
 ),
 page('Catálogo de requisitos de acesso e CRM',
   p('Local identifica requisito do piloto; Local parcial identifica cobertura incompleta descrita na secção Cobertura da versão local. Produção identifica desenvolvimento posterior. A rastreabilidade liga requisitos às histórias; a aceitação depende dos critérios e evidência de teste.'),
   table(['ID','Área','Requisito','Fase'],[[a,c,d] for a,_,c,d,_ in requirements[:11]],[51,0,339,78])
 ),
 page('Catálogo de requisitos de transporte e parque',
   table(['ID','Requisito','Fase'],[[a,c,d] for a,_,c,d,_ in requirements[11:22]],[51,339,78]),
   h('Validação do transporte'),
   p('Verificar sempre a pessoa responsável pelo serviço, o estado atual e as ligações a veículo e cliente. O campo pago não pode ser alterado livremente: deve ser calculado a partir dos recebimentos do documento interno. Os registos migrados preservam o estado antigo como informação a reconciliar.')
 ),
 page('Catálogo de requisitos de oficina e financeiro',
   table(['ID','Requisito','Fase'],[[a,c,d] for a,_,c,d,_ in requirements[22:31]],[51,339,78]),
   h('Validação financeira'),
   p('A emissão e o recebimento são eventos diferentes. Um documento pode estar emitido e não estar pago. O parcial mantém Emitida no documento e mostra Parcial no serviço; saldo nulo determina Paga e Pago. A soma não pode ultrapassar o valor devido. A referência impede a contabilização repetida da mesma evidência.')
 ),
 page('Catálogo de requisitos de entrega e produção',
   table(['ID','Requisito','Fase'],[[a,c,d] for a,_,c,d,_ in requirements[31:]],[51,339,78]),
   h('Critérios da decisão de produção'),
   b('Sergiu valida o processo completo e as responsabilidades; cada perfil executa o respetivo procedimento.','Os testes de autorização e de integridade são aprovados e não existem defeitos críticos conhecidos.','Restauro demonstrado e monitorização operacional atribuída.','Integração de faturação e contabilidade validada com o contabilista; âmbito de dados pessoais e conservação revisto.','Migração e eventual regresso ao sistema anterior ensaiados com cópia preservada.'),
   p('A decisão de produção deve registar o responsável, a versão, as evidências e os riscos aceites. A existência de formulários ou autenticação local, por si só, não prova segurança para acesso pela Internet.')
 ),
 page('Aceitação indicadores e referências',
   h('Cenários mínimos de aceitação'),
   b('Cliente > cotação > aprovação > aceitação > serviço > despacho > entrega > minuta > aprovação > recebimento parcial > recebimento integral.','Motorista tenta consultar serviço alheio; comercial tenta aprovar; administrador tenta decidir orçamento. Todas as ações indevidas são recusadas.','O requerente tenta aprovar o próprio registo; a regra normal recusa e a exceção explícita do proprietário exige motivo válido.','Duas atribuições com recurso indisponível, despacho de veículo em oficina, documento sem prova e recebimento superior ao saldo são recusados.','Repetição de conversão, saída e referência de pagamento não duplica os eventos.','Cópia e restauro preservam contas, relações e histórico; dados migrados permanecem disponíveis.'),
   h('Indicadores propostos'),
   table(['Indicador','Definição'],[
    ['Pontualidade','Entregas no prazo divididas pelas entregas com data prevista e efetiva.'],
    ['Saldo por cliente','Soma dos valores internos emitidos menos recebimentos reconciliados.'],
    ['Ocupação do parque','Posições ocupadas divididas pelas posições operacionais configuradas.'],
    ['Custo de oficina','Soma dos custos efetivos registados por veículo e período.'],
    ['Tempo de aprovação','Diferença entre submissão e decisão; apresentar os pendentes à parte.']
   ],[140,328]),
   h('Referências técnicas'),
   p('As referências sustentam os requisitos de segurança propostos. Não constituem certificação do produto. Consultadas em 30 setembro 2026.'),
   p('[1] OWASP Authentication Cheat Sheet  https://cheatsheetseries.owasp.org/cheatsheets/Authentication_Cheat_Sheet.html'),
   p('[2] OWASP Authorization Cheat Sheet  https://cheatsheetseries.owasp.org/cheatsheets/Authorization_Cheat_Sheet.html'),
   p('[3] OWASP Password Storage Cheat Sheet  https://cheatsheetseries.owasp.org/cheatsheets/Password_Storage_Cheat_Sheet.html')
 )
]

# Fix the first catalogue to the same three-column schema as the others.
PAGES['01_Requisitos_e_governacao'][6]['items'][-1] = table(['ID','Requisito','Fase'],[[a,c,d] for a,_,c,d,_ in requirements[:11]],[51,329,88])
PAGES['01_Requisitos_e_governacao'].insert(6,page('Cobertura da versão local',
 p('O catálogo define o objetivo do produto. Esta tabela descreve o alcance efetivo do piloto e os campos ou controlos que ainda exigem desenvolvimento. A validação de Sergiu deve usar ambos para evitar assumir que uma extensão prevista já está disponível.'),
 table(['Área','Disponível no piloto','Extensão necessária'],[
  ['Acessos','Uma conta inicial de Sergiu; um perfil por conta; controlo no servidor.','MFA, recuperação segura, identidade externa e revisão operacional de acesso.'],
  ['Cliente','Nome, NIF, contactos, responsável e limite aprovado informativo.','Moradas, estado, oportunidades e controlo automático de exposição ao crédito.'],
  ['Cotação','Percurso, datas, preço, aprovação, aceitação e conversão única.','Condições, anexos, propostas formais e envio externo.'],
  ['Planeamento','Motorista, conta associada opcional e matrícula; bloqueio no despacho e início.','Cadastro completo de recursos e janelas horárias; conflitos ao planear.'],
  ['Parque','Presença ligada ao serviço, posição livre em texto e movimentos.','Campo de portão, catálogo de posições, reservas e capacidade.'],
  ['Oficina','Matrícula, descrição, estimativa, estados, custo final, notas de conclusão e libertação.','Prioridade, peças, mão de obra e anexos de verificação técnica.'],
  ['Financeiro','Minuta não fiscal, vencimento, aprovação, emissão interna e recebimentos.','Faturação real, notas de crédito, estornos e reconciliação bancária.'],
  ['Continuidade','Base local, cópia de migração e exportação autorizada.','Restauro operacional repetível, cópias externas automatizadas e base transacional.']
 ],[83,193,192]),
 h('Sequência entre parque e transporte'),
 p('Operações planeia; parque confirma e regista saída física; operações autoriza o despacho no TMS; motorista inicia o serviço. O sistema recusa despacho ou início enquanto existe presença ativa no parque. A saída do parque também é recusada se existir bloqueio de oficina.'),
 h('Evidência fora do formulário'),
 p('Quando falta um campo, conservar informação em apoio restrito autorizado por Sergiu. Um nome de motorista em texto não concede acesso individual; associar a conta ao serviço. O limite de crédito aprovado é informativo: o piloto não calcula exposição nem bloqueia automaticamente despacho ou emissão por ultrapassagem.')
))

def sop(code, title, trigger, responsible, pre, steps, approvals, evidence, exception, req, target=''):
    items = [p('Procedimento '+code+'   Responsável '+responsible),h('Quando executar'),p(trigger),h('Condições de início'),p(pre),h('Passos')]
    items += [p(str(i+1)+'. '+s) for i,s in enumerate(steps)]
    items += [h('Aprovação e evidência'),p(approvals),p(evidence),h('Exceção e escalonamento'),p(exception),p('Rastreabilidade  '+req)]
    if target: items += [p('Extensão de produção  '+target)]
    return page(code+' '+title,*items)

PAGES['02_Procedimentos_operacionais'] = [
 page('Procedimentos operacionais da plataforma TSS',
  p('Este manual orienta o trabalho diário desde o registo do cliente até ao recebimento, incluindo parque, oficina, acessos e recuperação. Cada procedimento identifica quem executa, o que tem de existir, como avançar, quem aprova e que evidência conservar.'),
  table(['Controlo','Valor'],[['Empresa',COMPANY],['Responsável','Sergiu Sandu'],['Versão e data','1.0   '+DATE],['Âmbito','Piloto local e preparação dos procedimentos de produção']],[140,328]),
  h('Como utilizar'),
  p('Executar o procedimento correspondente à tarefa. O identificador SOP liga os passos aos requisitos e às histórias. Quando a interface não suporta ainda uma extensão de produção, guardar a evidência na localização restrita aprovada por Sergiu e manter a referência no registo.'),
  h('Regras comuns'),
  b('Confirmar o cliente, veículo ou serviço antes de gravar qualquer ação.','Usar conta própria e perfil autorizado; não partilhar palavras passe.','Guardar referências que permitam localizar documentos e decisões.','Gerente aprova todas as decisões indicadas; não existem escalões monetários neste piloto.','A autoaprovação normal é recusada. Sergiu sozinho pode usar a exceção explícita do proprietário, com motivo de pelo menos 20 caracteres.','Uma minuta interna não é documento fiscal. A faturação real continua no sistema validado pela empresa.'),
  h('Responsabilidades no ensaio'),
  p('Sergiu Sandu é o gerente e o único titular nominal confirmado. Comercial, operações, motorista, parque, oficina, financeiro, auditor e administrador são funções a atribuir. Sergiu pode testar as funções durante a validação; as atribuições futuras devem constar da lista de utilizadores.'),
  p('Antes de utilizar dados reais, confirmar cópia de segurança, acesso físico ao computador e localização de provas e documentos. Definir o responsável de cada turno e a forma de comunicar incidentes sem expor informação de clientes.')
 ),
 sop('SOP01','Registar cliente e rever crédito',
  'Surge um novo pedido de transporte, um novo cliente ou uma alteração às condições de pagamento.','Comercial com financeiro; gerente decide crédito.',
  'Existe identificação do cliente e contacto de referência. A pessoa tem acesso autorizado ao CRM.',
  ['Pesquisar nome e identificação antes de criar a ficha. Confirmar se o cliente já existe.','Registar nome, identificação e contactos. O piloto tem contacto, email e telefone; guardar moradas e informação adicional no registo de apoio restrito.','Registar referência do pedido e condições de pagamento pretendidas no apoio. Financeiro verifica saldo e condições com comercial.','Comercial prepara pedido de alteração de limite, com valor atual, proposto e justificação, e submete ao gerente. Financeiro fornece a informação de suporte.','Gerente revê a informação, decide e regista o motivo. Comunicar internamente a decisão ao comercial.','Só usar o limite aprovado. Se houver dúvidas sobre identificação ou capacidade de pagamento, manter o pedido em análise.'],
  'A alteração de crédito exige aprovação de Sergiu. A exceção do proprietário aplica-se apenas à separação de pessoas e fica justificada.',
  'Conservar ficha, referência de contacto, pedido de crédito, decisão e histórico de alterações.',
  'Duplicação ou divergência de identificação é escalada ao gerente. Não apagar fichas ligadas a serviços. O limite é informativo: não existe bloqueio automático de despacho ou emissão por exposição. Limite zero não impõe pagamento antecipado; acordar condições concretas.',
  'REQ07 REQ09   US07 US09',
  'Definir verificações de cliente e política de crédito, exposição comprometida e bloqueio automático quando houver condições aprovadas. O piloto não substitui análise de risco.'),
 sop('SOP02','Preparar aprovar e aceitar cotação',
  'O cliente pede preço e disponibilidade para um transporte.','Comercial; gerente aprova.',
  'Cliente identificado. Conhecem-se locais de recolha e entrega, datas e necessidades do serviço.',
  ['Criar cotação em Rascunho ligada ao cliente. Registar recolha, entrega, datas e preço em euros.','Verificar percurso, disponibilidade preliminar, condições e consistência do valor. Guardar pressupostos no registo de apoio restrito, pois o piloto não tem campo de condições.','Submeter a cotação em Em aprovação. Aguardar decisão; não marcar aceitação de cotação ainda não aprovada.','Gerente revê informação, aprova ou rejeita com motivo. Corrigir o registo rejeitado e voltar a submeter quando permitido.','Registar a aceitação do cliente somente após aprovação, incluindo referência verificável de mensagem, proposta assinada ou outro acordo.','Passar a cotação aceite a operações para conversão. Confirmar que a mesma cotação não criou já um serviço.'],
  'Sergiu aprova a libertação comercial. No ensaio individual, seleciona a exceção do proprietário e explica a acumulação de funções.',
  'Conservar versão da cotação, decisão, referência da aceitação e ligação ao serviço criado.',
  'Mudança de preço ou de percurso após aprovação exige revisão e nova decisão antes do compromisso. Se o cliente não aceitar, conservar a cotação e o motivo. O envio externo não é automático no piloto.',
  'REQ10 REQ11   US10 US11 US12'),
 sop('SOP03','Planear atribuir e despachar serviço',
  'Uma cotação foi aceite ou operações recebeu um serviço autorizado para planear.','Operações.',
  'Cliente e datas válidos, motorista identificado, matrícula conhecida e acesso autorizado ao TMS.',
  ['Converter uma vez a cotação aceite em serviço Agendado. Confirmar referência, cliente, percurso e preço.','Atribuir motorista e matrícula adequados. Associar conta motorista para acesso individual. Rever conflitos; o piloto valida dias completos no despacho e início.','Confirmar disponibilidade, locais e instruções. Verificar documentação e condições de condução pelos procedimentos da empresa; estes controlos humanos não são automatizados.','Se existe presença no parque, solicitar conferência e saída pelo SOP05 antes do despacho. Confirmar que não há bloqueio de oficina.','Autorizar despacho no TMS. Confirmar hora gravada e recursos. O sistema recusa presença ativa no parque ou recurso em conflito.','O motorista inicia a execução e operações acompanha o estado. Depois do despacho, o planeamento fica bloqueado no piloto.'],
  'A afetação de recursos pertence a operações. Alterações comerciais exigem revisão do gerente; o despacho não substitui essa decisão.',
  'Conservar serviço, afetação, despacho, instruções e verificação operacional aplicável.',
  'Conflito de recurso ou bloqueio de oficina impede o despacho. Reatribuir recurso disponível ou escalar ao gerente; nunca apagar a ordem de oficina para desbloquear o veículo.',
  'REQ12 REQ13 REQ14   US13 US14 US15',
  'Integrar disponibilidade, qualificações, documentação e telemática. O piloto valida conflitos e estados, não comprova todas as condições legais de operação.'),
 sop('SOP04','Controlar entrada e movimento no parque',
  'Um veículo chega ao parque ou tem de mudar de posição.','Parque.',
  'Veículo identificado, portão e posição conhecidos; existe acesso ao YMS.',
  ['Confirmar matrícula, serviço associado e presença física. Consultar se já existe presença ativa para evitar entrada duplicada.','Registar entrada com veículo, serviço e posição disponível. O instante é gravado. Referir o portão na observação, pois ainda não existe campo próprio.','Conferir a posição no terreno e registar observações sobre condição ou necessidade de apoio.','Para movimentar, confirmar origem e destino livre. Registar movimento antes de considerar a mudança concluída no sistema.','Confirmar que a nova posição corresponde à localização real. Corrigir divergência através de ação rastreável.','Comunicar bloqueios de circulação ou avarias a operações e oficina; quando apropriado, abrir ordem de trabalho.'],
  'O perfil parque executa movimentos normais. Uma exceção de segurança ou indisponibilidade deve ser decidida pelo responsável operacional e escalada a Sergiu quando altera o serviço.',
  'Conservar entrada, portão, posições, instantes, ator e observações do movimento.',
  'Matrícula desconhecida, posição ocupada ou presença duplicada exige conferência física antes de registar. Não criar uma segunda matrícula para contornar a validação. Em falha de sistema, usar registo de contingência e reconciliar depois.',
  'REQ19 REQ20   US21 US22',
  'Configurar capacidade, reservas, mapa de posições e ligação a portarias e sensores. Não é autorizada movimentação física pela simples existência de um registo.'),
 sop('SOP05','Controlar saída do parque',
  'Operações solicita saída ou o veículo apresenta-se no portão.','Parque com operações.',
  'Existe presença ativa e veículo identificado; existe destino ou serviço associado à saída.',
  ['Confirmar matrícula, presença ativa, serviço associado, motorista e motivo da saída.','Verificar o serviço previsto e consultar bloqueios de oficina. Um veículo bloqueado não pode ter saída operacional.','Confirmar com operações a instrução de saída. A autorização de despacho no TMS ocorre depois do registo de saída do parque.','Registar saída com instante e observação de referência. A posição deve ficar disponível.','Conferir que o evento foi gravado uma vez e que deixou de existir presença ativa.','Informar operações para autorizar o despacho no TMS. Conservar referência de qualquer divergência.'],
  'Parque regista a saída normal autorizada. Só Sergiu pode libertar um veículo após conclusão técnica da oficina. Uma saída para assistência externa necessita de decisão e procedimento específicos, fora do fluxo normal do piloto.',
  'Conservar evento de saída, portão ou referência, serviço relacionado e identidade do operador.',
  'Sem presença ativa, não inventar uma entrada retroativa sem verificação. Conferir movimentos reais e escalar divergências. Um bloqueio técnico deve ser resolvido pela oficina e pelo gerente.',
  'REQ21   US23'),
 sop('SOP06','Registar entrega e tratar incidente',
  'A recolha ou entrega terminou, ou surgiu avaria, atraso, dano ou impossibilidade de entrega.','Motorista do serviço; operações acompanha.',
  'Serviço atribuído ao motorista e em estado compatível com a execução.',
  ['Consultar o serviço atribuído e confirmar cliente, destino e referência.','Se houver incidente, contactar operações pelo canal aprovado e registar momento, localização, descrição e evidência disponível.','Na entrega, confirmar destinatário e documento ou referência da prova. Registar esses dados de forma legível.','Registar a conclusão somente quando existe prova de entrega. Operações confere a coerência com o serviço.','Guardar a prova na localização autorizada e a referência no registo; não usar armazenamento público.','Quando a entrega falha, manter o serviço em acompanhamento, acordar nova instrução e conservar o histórico.'],
  'A conclusão operacional depende de prova. Alterações de preço, novo compromisso ou cancelamento são decididos pelo gerente nos fluxos apropriados.',
  'Conservar destinatário, referência da prova, serviço, momento e qualquer decisão sobre incidente.',
  'Não concluir para ultrapassar um bloqueio de faturação. Se um serviço importado já estiver concluído sem prova, adicionar e rever a prova antes da minuta. Dano, emergência ou risco pessoal segue primeiro os procedimentos de segurança da empresa.',
  'REQ15 REQ16   US16 US17 US18',
  'Adicionar anexos protegidos, registo estruturado de incidentes, confirmação de receção e aplicação móvel quando o âmbito estiver aprovado.'),
 sop('SOP07','Pedir e decidir cancelamento',
  'O cliente retira o pedido ou o serviço deixa de ser executável.','Operações pede; gerente decide.',
  'Serviço identificado e motivo conhecido. Custos ou compromissos existentes foram conferidos.',
  ['Consultar estado, recursos, cotação e documentos ligados ao serviço.','Registar pedido de cancelamento com motivo e referência da instrução recebida.','Indicar o tratamento proposto de custos, comunicação ao cliente e recursos afetados em observações.','Gerente verifica e aprova ou rejeita o pedido com motivo.','Após aprovação, confirmar estado Cancelado e disponibilidade dos recursos; a presença de parque e a ordem de oficina continuam a exigir tratamento próprio.','Financeiro verifica se existem valores internos ou recebimentos a corrigir por procedimento controlado.'],
  'Sergiu decide o cancelamento. Se for também requerente no ensaio, usa a exceção explícita do proprietário com justificação válida.',
  'Conservar pedido, decisão, comunicação de referência e ligação aos documentos afetados.',
  'Um serviço com documento emitido ou recebimento não deve ser apagado. Escalar a correção ao gerente e ao financeiro; notas de crédito e devoluções fiscais pertencem ao sistema de faturação externo.',
  'REQ17   US19'),
 sop('SOP08','Abrir aprovar e executar ordem de oficina',
  'Existe avaria, revisão ou reparação a realizar num veículo.','Oficina; gerente aprova o orçamento.',
  'Veículo identificado, descrição do problema e estimativa disponíveis.',
  ['Abrir ordem em Rascunho com matrícula, problema e estimativa. Indicar prioridade na descrição, pois não existe campo próprio. Conferir bloqueio operacional.','Informar operações para rever serviços atribuídos. O piloto exige que o veículo não tenha serviço Em curso para abrir intervenção.','Submeter orçamento a aprovação com trabalhos previstos na descrição e motivo de suporte.','Sergiu aprova ou rejeita. Oficina só inicia trabalhos na ordem aprovada.','Registar custo efetivo e notas de conclusão com resumo do trabalho. Conservar verificação técnica em apoio restrito e respetiva referência, pois não existem anexos próprios.','Marcar Concluída e solicitar libertação ao gerente; não considerar o veículo disponível apenas porque terminou o trabalho.'],
  'Sergiu aprova orçamento e posterior libertação. Despesas fora do orçamento aprovado devem ser submetidas a nova decisão antes do compromisso, segundo procedimento da empresa.',
  'Conservar diagnóstico, estimativa, aprovação, trabalhos, custo efetivo e verificação técnica.',
  'Risco de segurança impede circulação. Divergência de custos ou trabalhos exige nova decisão; o piloto não substitui inspeção técnica nem controlo de qualificações. Uma ordem rejeitada conserva o histórico e o problema deve ter encaminhamento.',
  'REQ23 REQ24   US25 US26 US27',
  'Implementar peças, mão de obra, fornecedores e revisão formal do orçamento. O custo final não prova por si só a aprovação de uma alteração.'),
 sop('SOP09','Rever e libertar veículo da oficina',
  'A oficina concluiu o trabalho e pede devolução do veículo às operações.','Gerente; oficina fornece evidência.',
  'Ordem Concluída com custo efetivo e descrição do trabalho; verificação técnica disponível.',
  ['Consultar ordem, diagnóstico, aprovação inicial e conclusão técnica.','Confirmar que o trabalho necessário foi realizado e que não existem bloqueios técnicos pendentes.','Conferir custo efetivo e alterações. Se excede a estimativa, a libertação exige motivo com pelo menos 20 caracteres para aprovar o custo final.','Sergiu regista a decisão de libertação com motivo. Normalmente deve ser pessoa diferente de quem concluiu.','Confirmar estado Libertada e disponibilidade operacional do veículo. Rever outras ordens ativas antes de despachar.','Informar operações e parque; atualizar planeamento e guardar a evidência da inspeção.'],
  'Só gerente liberta. No ensaio individual, a exceção do proprietário identifica claramente que Sergiu acumulou conclusão e decisão e grava a justificação.',
  'Conservar identidade de quem concluiu, decisão de libertação, motivo, instante e referência técnica.',
  'Sem evidência ou com dúvida técnica, manter bloqueio e solicitar verificação competente. A aprovação no software não constitui certificação de segurança do veículo.',
  'REQ25   US28'),
 sop('SOP10','Preparar e aprovar documento interno',
  'O transporte está concluído e existe prova de entrega para preparar cobrança.','Financeiro; gerente aprova.',
  'Serviço Concluído, cliente correto, prova com destinatário e referência, preço conferido.',
  ['Consultar o serviço e confirmar que não existe já documento interno ligado.','Criar a minuta com valor do serviço, referência e data de vencimento acordada. Conferir prova e dados do cliente.','Submeter a minuta em Em aprovação; não registar recebimento antes da etapa permitida.','Sergiu revê e aprova ou rejeita com justificação. Corrigir divergências mantendo histórico.','Após aprovação, financeiro marca emissão interna. Confirmar que o ecrã identifica o documento como interno e não fiscal.','Para operação real, emitir através do sistema de faturação da empresa e manter a referência externa quando a integração estiver disponível.'],
  'Sergiu aprova a minuta. A exceção do proprietário exige motivo quando o próprio preparou o registo no ensaio.',
  'Conservar serviço, cotação de origem, prova, minuta, decisão e referência da emissão real quando aplicável.',
  'Preço incorreto ou falta de prova exige corrigir a origem e repetir a revisão. Não emitir documento fiscal a partir deste piloto. Faturação, impostos, notas de crédito e requisitos portugueses exigem sistema e validação próprios.',
  'REQ27 REQ28 REQ31   US30 US31 US34'),
 sop('SOP11','Registar recebimento e acompanhar saldo',
  'Existe evidência de recebimento ou um valor vencido a acompanhar.','Financeiro.',
  'Documento interno Emitida, valor e saldo conferidos; existe referência verificável do pagamento.',
  ['Localizar cliente e documento. Comparar evidência, montante, data e referência antes de registar.','Introduzir montante positivo não superior ao saldo e referência que ainda não esteja usada.','Gravar recebimento e confirmar saldo atualizado. Um recebimento parcial mantém o documento Emitida; saldo nulo muda para Paga.','Reconciliar com a evidência bancária ou caixa autorizada. Registo no piloto não movimenta dinheiro nem confirma automaticamente o extrato.','Rever documentos vencidos com saldo e registar referência do contacto de cobrança na localização aprovada.','Escalar discrepâncias ou pedidos de devolução ao gerente e conservar a decisão antes da correção financeira.'],
  'Financeiro reconcilia recebimentos. Alterações de crédito ou exceções comerciais pertencem ao gerente. Devoluções e ajustamentos precisam de procedimento próprio e não devem ser simulados com montantes negativos.',
  'Conservar valor, referência, data, documento, ator e saldo; ligar a evidência do extrato sem expor dados desnecessários.',
  'Valor superior ao saldo, referência repetida ou cliente incorreto é recusado e revisto. Pagamentos históricos migrados devem ser reconciliados, evitando registar novamente valores já recebidos.',
  'REQ29 REQ30   US32 US33',
  'Implementar importação de extrato, propostas de correspondência e autorização de estornos antes de uso financeiro em escala.'),
 sop('SOP12','Administrar utilizadores e responsabilidades',
  'Uma pessoa entra, muda de responsabilidade, deixa a empresa ou necessita de acesso revisto.','Administrador executa; gerente autoriza funções.',
  'Existe decisão de Sergiu, identidade da pessoa e funções necessárias; administrador autenticado.',
  ['Registar pedido e decisão sobre acesso. Confirmar pessoa e necessidade; não criar contas partilhadas.','Criar conta nominal e atribuir um perfil aprovado. Associar a conta motorista aos serviços aplicáveis. Só gerente gere contas de gerente.','Entregar credencial pelo canal aprovado e orientar a pessoa para alterar a palavra passe. Não guardar a palavra passe no procedimento.','Testar acesso permitido e recusa de tarefa indevida. Verificar que administrador não recebe poder comercial por omissão.','Quando muda a função, rever perfil, sessões e responsabilidade de registos pendentes.','Na saída, desativar conta, transferir tarefas e conferir recusa de acesso. Manter histórico da pessoa.'],
  'Sergiu autoriza atribuições e acumulações; administrador executa tecnicamente. O piloto deve conservar a referência da decisão em observações ou histórico.',
  'Conservar pedido, decisão, pessoa, perfis, associação, data de ativação ou desativação e testes de acesso.',
  'Não remover o último gerente ativo sem sucessão aprovada. Suspeita de acesso indevido exige desativação ou revogação pelo procedimento disponível e escalonamento imediato. Recuperação de conta não deve revelar a palavra passe antiga.',
  'REQ01 REQ02 REQ03 REQ06   US01 US02 US03 US04 US38 US39',
  'Adicionar convite com prazo, MFA, recuperação segura e revogação central de sessões. Revisões periódicas ficam registadas.'),
 sop('SOP13','Criar cópia e testar recuperação',
  'Antes de migração ou atualização, segundo calendário de cópias, ou quando há perda ou corrupção.','Gerente exporta; administrador executa a recuperação autorizada.',
  'Existe acesso autorizado à cópia completa e localização restrita. Conhecem-se versão da aplicação e momento da cópia.',
  ['Criar cópia pelo mecanismo autorizado. Para cópia manual, parar a aplicação antes de copiar a base completa e reiniciar depois.','Registar data, versão, ficheiro, localização, responsável e resultado. Conservar a cópia anterior à migração.','Guardar cópia fora da pasta operacional e restringir acesso; contas e histórico também contêm informação sensível.','Testar numa cópia separada da aplicação. Parar o serviço, preservar a base atual e colocar a cópia de teste no caminho indicado no guia técnico.','Iniciar a cópia de teste e conferir contas, clientes, serviços, decisões, provas, minutas, recebimentos e saldo. Não testar por cima da única base de trabalho.','Sergiu aceita o resultado ou decide recuperação real. Registar período eventualmente perdido e reconciliar operações efetuadas depois da cópia.'],
  'A recuperação real exige decisão de Sergiu. O administrador verifica integridade antes de reabrir acesso e conserva a base anterior à recuperação.',
  'Conservar cópia, registo do ensaio, verificações, decisão e tempo medido. A cópia no mesmo computador não protege por si só contra perda desse computador.',
  'Cópia incompleta ou falha de acesso impede recuperação. Manter evidência, suspender gravações e escalar. Não editar diretamente o JSON para corrigir saldos sem diagnóstico e plano de recuperação.',
  'REQ32 REQ33 REQ34   US35 US36 US37',
  'Automatizar cópias cifradas, política de conservação, recuperação transacional e monitorização. Validar objetivos de perda de dados e tempo de recuperação.'),
 page('Modelos de registo e revisão diária',
  h('Registo de aprovação'),
  p('Entidade e referência  ____________________   Versão  __________\nRequerente  ____________________   Data do pedido  __________\nDecisor  ____________________   Data da decisão  __________\nResultado  Aprovado ou Rejeitado\nMotivo  _____________________________________________________\nExceção explícita do proprietário  Sim ou Não\nJustificação da acumulação de funções  _________________________'),
  h('Registo de incidente'),
  p('Serviço ou veículo  ____________________   Momento  __________\nLocal e descrição  ____________________________________________\nResponsável de acompanhamento  ______________________________\nEvidência de referência  _______________________________________\nDecisão e instrução seguinte  __________________________________\nResultado e momento de fecho  _________________________________'),
  h('Registo de cópia e recuperação'),
  p('Versão e data da cópia  _______________________________________\nFicheiro e localização restrita  _________________________________\nOperador  ____________________   Resultado  __________________\nConta e acessos verificados  ___________________________________\nRelações e saldos verificados  __________________________________\nTempo de recuperação medido  _________________________________\nDecisão do gerente e referência  _______________________________'),
  h('Revisão no fim do dia'),
  b('Operações revê serviços sem atualização, recursos indisponíveis e entregas sem prova.','Parque compara presenças e posições com a realidade.','Oficina revê ordens bloqueadas, aprovações e pedidos de libertação.','Financeiro revê minutas pendentes, referências de pagamentos e saldos.','Gerente revê decisões pendentes e exceções do proprietário.','Administrador confirma cópia, falhas e acessos que exigem revisão.')
 )
]

stories = []
def us(id, module, role, goal, benefit, criteria, priority, sprint, req, dep='', phase='Local'):
    stories.append(dict(id=id,module=module,role=role,goal=goal,benefit=benefit,criteria=criteria,priority=priority,sprint=sprint,requirements=req.split(),dependency=dep,phase=phase))

us('US01','Acessos','gerente','criar a minha conta inicial sem contas partilhadas','identificar cada ação',['A primeira configuração cria apenas Sergiu com credencial escolhida.','Novo arranque mantém a conta e não repete a configuração.'],'P0',1,'REQ01')
us('US02','Acessos','utilizador','entrar sair e mudar a palavra passe','proteger o acesso',['Credencial errada é recusada; palavra passe não fica em texto simples.','Saída e expiração impedem novos pedidos com sessão inválida.'],'P0',1,'REQ02','US01')
us('US03','Acessos','administrador','atribuir perfis e desativar contas','manter responsabilidades atuais',['Perfis válidos são guardados e usados pelo servidor.','Conta desativada perde acesso e conserva histórico.'],'P0',1,'REQ03','US01')
us('US04','Acessos','motorista','consultar apenas os meus serviços','proteger dados de outros serviços',['Serviço atribuído é acessível no ecrã e na API.','Identificador de serviço alheio é recusado e não aparece no CSV.'],'P0',1,'REQ03','US03')
us('US05','Aprovações','gerente','decidir pedidos com separação de funções','tornar a decisão verificável',['Autoaprovação normal é recusada.','Exceção explícita do proprietário exige 20 caracteres e grava requerente, ator, data e motivo.'],'P0',2,'REQ04','US03')
us('US06','Histórico','auditor','consultar decisões e alterações','seguir a origem de um registo',['Ação autorizada gera evento ligado à entidade e ao ator.','Auditor consulta histórico e não altera registos.'],'P0',2,'REQ05','US03')
us('US07','CRM','comercial','registar uma ficha de cliente','reutilizar os dados comerciais',['Campos essenciais validados e pesquisa antes da criação.','Cliente mantém identificador estável em cotações e serviços.'],'P0',2,'REQ07','US03')
us('US08','CRM','comercial','acompanhar oportunidades e contactos','saber o próximo passo comercial',['Cada oportunidade tem etapa, responsável e tarefa seguinte.','Contacto fica ligado ao cliente com data e resultado.'],'P1',2,'REQ08','US07','Produção')
us('US09','CRM','comercial','pedir alteração de limite de crédito','usar condições aprovadas',['Pedido conserva limite anterior, proposto e justificação com apoio financeiro.','Só gerente decide e o novo valor aparece após aprovação.'],'P0',3,'REQ09','US05 US07')
us('US10','CRM','comercial','preparar cotação com percurso datas e preço','propor um serviço definido',['Cliente, locais, datas e valor em euros são obrigatórios.','Rascunho permite revisão sem marcar aceitação.'],'P0',3,'REQ10','US07')
us('US11','CRM','gerente','aprovar ou rejeitar uma cotação','autorizar o compromisso comercial',['Só cotação Em aprovação admite decisão.','Decisão grava motivo e segue a regra de autoaprovação.'],'P0',3,'REQ10','US05 US10')
us('US12','CRM TMS','operacoes','converter cotação aceite em serviço','manter origem comercial',['Aceitação exige cotação aprovada e referência do cliente.','Uma conversão cria um serviço e repetição não duplica.'],'P0',3,'REQ11','US11')
us('US13','TMS','operacoes','manter ficha completa de serviço','acompanhar a execução',['Cliente, locais, datas, valor e estado são validados.','Os recursos e a cotação de origem permanecem ligados.'],'P0',4,'REQ12','US12')
us('US14','TMS','operacoes','atribuir veículo e motorista disponíveis','evitar conflitos',['Piloto recusa sobreposição de dias completos no despacho e início; produção valida também no planeamento.','Veículo com ordem ativa de oficina fica indisponível.'],'P0',4,'REQ13','US13 US25','Local e produção')
us('US15','TMS','operacoes','despachar o serviço pronto','registar autorização operacional',['Sem motorista, matrícula, estado válido ou com presença ativa no parque o despacho é recusado.','Hora do despacho é gravada uma vez sem iniciar entrega.'],'P0',4,'REQ14','US14')
us('US16','TMS','motorista','iniciar e concluir o serviço atribuído','manter estado coerente',['Só transições permitidas são aceites.','Conclusão exige prova com destinatário e referência.'],'P0',4,'REQ15','US04 US15')
us('US17','TMS','operacoes','rever a prova de entrega','preparar cobrança com evidência',['Prova fica ligada ao serviço correto.','Serviço histórico concluído sem prova permite registo e revisão antes da minuta.'],'P0',5,'REQ15','US16')
us('US18','TMS','motorista','registar incidentes com evidência','obter uma instrução acompanhada',['Tipo, momento e descrição são obrigatórios.','Operações atribui acompanhamento e decisão com histórico.'],'P1',5,'REQ16','US16','Produção')
us('US19','TMS','gerente','decidir um cancelamento pedido','preservar compromissos e histórico',['Operações pede com motivo e gerente decide.','Ação repetida não altera duas vezes nem apaga relações financeiras.'],'P0',5,'REQ17','US05 US13')
us('US20','TMS','utilizador','pesquisar filtrar e exportar serviços','obter uma lista útil',['Filtros combinados limitam os resultados visíveis.','CSV preserva euros, acentos e apenas registos autorizados.'],'P0',5,'REQ18','US03 US13')
us('US21','YMS','parque','registar entrada numa posição livre','saber que veículos estão no parque',['Uma presença ativa por veículo e posição exclusiva, ligada ao serviço.','Entrada grava instante e operador; portão fica em observação no piloto e em campo próprio na produção.'],'P0',6,'REQ19','US03','Local e produção')
us('US22','YMS','parque','mover presença para nova posição','manter localização correta',['Destino ocupado é recusado.','Movimento conserva origem, destino, ator e momento.'],'P0',6,'REQ20','US21')
us('US23','YMS','parque','registar saída autorizada','libertar a posição corretamente',['Sem presença ativa ou com bloqueio operacional a saída é recusada.','Saída termina presença e repetição não cria outro evento.'],'P0',6,'REQ21','US21 US25')
us('US24','YMS','operacoes','configurar capacidade e reservas do parque','planear chegada e espaço',['Posições operacionais e capacidade são configuráveis.','Reserva em conflito produz indicação e impede sobreocupação.'],'P1',6,'REQ22','US21','Produção')
us('US25','Oficina','oficina','abrir ordem ligada ao veículo','tornar visível a indisponibilidade',['Problema e estimativa são validados.','Ordem ativa bloqueia despacho até rejeição ou libertação.'],'P0',7,'REQ23','US03')
us('US26','Oficina','gerente','aprovar o orçamento da ordem','autorizar os trabalhos',['Oficina submete; gerente aprova ou rejeita.','Início é recusado antes de aprovação válida.'],'P0',7,'REQ24','US05 US25')
us('US27','Oficina','oficina','registar execução e conclusão','demonstrar o trabalho realizado',['Ordem aprovada permite início e conclusão.','Conclusão exige custo efetivo, notas do trabalho e responsável rastreável.'],'P0',7,'REQ24','US26')
us('US28','Oficina','gerente','libertar veículo após conclusão','devolver disponibilidade com decisão',['Só ordem Concluída permite libertação pelo gerente.','Quem concluiu não decide sem exceção explícita válida.'],'P0',7,'REQ25','US05 US27')
us('US29','Oficina','oficina','acompanhar manutenção e peças','preparar revisões e stock',['Plano cria vencimento por data ou quilometragem.','Consumo de peça altera stock e liga custo ao veículo.'],'P1',8,'REQ26','US25','Produção')
us('US30','Financeiro','financeiro','criar minuta ligada ao serviço concluído','preparar cobrança rastreável',['Sem conclusão ou prova a criação é recusada.','Uma minuta por serviço e valor igual ao preço válido.'],'P0',8,'REQ27','US17')
us('US31','Financeiro','gerente','aprovar a minuta antes da emissão interna','controlar o valor a cobrar',['Financeiro submete e gerente decide com motivo.','Emissão interna só aceita aprovada e identifica caráter não fiscal.'],'P0',8,'REQ28','US05 US30')
us('US32','Financeiro','financeiro','registar pagamentos parciais e integrais','calcular saldo correto',['Valor positivo até ao saldo e referência única.','Parcial mantém Emitida e serviço Parcial; saldo zero determina Paga e serviço Pago.'],'P0',8,'REQ29','US31')
us('US33','Financeiro','financeiro','ver saldos e acompanhar vencidos','organizar cobranças',['Saldo deriva de recebimentos e não de campo livre.','Vencidos são identificados; extensão conserva contactos e próxima ação.'],'P0',9,'REQ30','US32','Local e produção')
us('US34','Financeiro','financeiro','ligar documento fiscal externo ao serviço','evitar dupla faturação real',['Integração usa referência externa e confirmação de emissão.','Falha e repetição não produzem emissão fiscal duplicada.'],'P0',9,'REQ31','US31','Produção')
us('US35','Dados','administrador','migrar serviços existentes com cópia','preservar trabalho anterior',['Cópia exata anterior à migração é preservada.','Serviços permanecem e estado antigo de pagamento exige reconciliação.'],'P0',9,'REQ32','US13')
us('US36','Continuidade','administrador','exportar e restaurar cópia completa','recuperar o sistema',['Exportação completa é restrita e inclui relações e contas.','Restauro de teste preserva acessos, histórico e saldos sem substituir a única base.'],'P0',9,'REQ33','US35','Local e produção')
us('US37','Qualidade','gerente','validar o ciclo completo e recusas','aceitar uma versão comprovada',['Ciclo completo e cenários negativos deixam evidência.','Defeito crítico impede decisão de produção.'],'P0',9,'REQ34','US19 US23 US28 US32 US36','Local e produção')
us('US38','Segurança','administrador','usar HTTPS MFA e monitorização','preparar acesso remoto seguro',['Perfis privilegiados usam MFA e transporte HTTPS.','Falhas, revogação de sessão e tentativas abusivas são testadas.'],'P0',10,'REQ06','US02 US03','Produção')
us('US39','Segurança','utilizador','recuperar acesso e confirmar ações sensíveis','manter a conta protegida',['Recuperação usa verificação, prazo e utilização única.','Mudança sensível revoga sessões conforme política e é auditada.'],'P0',10,'REQ06','US38','Produção')
us('US40','Dados pessoais','gerente','definir acesso conservação e pedidos','tratar informação de forma controlada',['Inventário identifica finalidade, acesso e prazo proposto.','Pedido de titular e eliminação controlada têm responsáveis e evidência.'],'P0',10,'REQ35','US03 US06','Produção')
us('US41','Plataforma','operacoes','gravar em simultâneo sem perda','trabalhar com equipa',['Duas alterações concorrentes produzem conflito controlado ou transação válida.','Unicidade de pagamento e conversão resiste à concorrência.'],'P0',10,'REQ36','US12 US32','Produção')
us('US42','Integrações','operacoes','receber eventos de serviços externos','evitar repetição manual',['Evento tem origem e identificador e não duplica ao repetir.','Falhas permitem reprocessamento e reconciliação por responsável.'],'P1',10,'REQ37','US41','Produção')
us('US43','Operação','gerente','treinar perfis e aceitar o piloto acompanhado','decidir entrada em produção',['Cada função executa SOP com evidência e resultado.','Gerente assina decisão ligada a versão, riscos e plano de recuperação.'],'P0',10,'REQ38','US34 US37 US38 US40 US41','Produção')
us('US44','Gestão','gerente','ver indicadores com definição e origem','acompanhar a atividade',['Período e população do indicador são explícitos.','Valor pode ser conciliado com registos autorizados de origem.'],'P1',10,'REQ39','US13 US21 US27 US32','Produção')

sprints = [
 (1,'Identidades e acesso','Configuração, entrada, perfis e alcance do motorista.','Demonstrar recusas na API e no ecrã.'),
 (2,'Decisões e relação comercial','Aprovações, histórico e cliente; desenho de oportunidades.','Sergiu revê responsabilidades e a exceção do proprietário.'),
 (3,'Crédito e cotação aceite','Pedido de crédito, cotação, aprovação, aceitação e conversão.','Uma cotação aceite gera um único serviço rastreável.'),
 (4,'Planeamento e execução','Recursos, despacho, estados e disponibilidade da oficina.','Resolver dependência de US25 com bloqueio mínimo antes de fechar US14.'),
 (5,'Prova incidentes e cancelamento','Prova, acompanhamento, cancelamento, pesquisa e exportação.','Sem prova não há conclusão elegível para minuta.'),
 (6,'Parque','Presenças, posições, movimentos, saída e capacidade.','Conferir ocupação real e recusar movimentos incompatíveis.'),
 (7,'Oficina','Ordens, aprovação, execução e libertação.','Despacho só retoma após tratamento de todos os bloqueios.'),
 (8,'Documentos e recebimentos','Minutas internas, aprovação, emissão interna e reconciliação.','Parcial e integral mantêm saldo correto; não há emissão fiscal.'),
 (9,'Migração recuperação e integração fiscal','Saldos, preservação de dados, restauro e desenho da integração.','A integração real depende do fornecedor escolhido e do contabilista.'),
 (10,'Preparação da produção','Segurança remota, dados, concorrência, integrações e aceitação.','Replanear por capacidade; não assumir que todas as integrações cabem num sprint.')
]

PAGES['03_Backlog_e_plano_de_sprints'] = [
 page('Backlog e plano de sprints da plataforma TSS',
  p('Este backlog transforma os requisitos em 44 histórias verificáveis para CRM, TMS, YMS, oficina, financeiro e plataforma. O plano propõe dez sprints de duas semanas. É uma sequência de trabalho a estimar com a equipa, sem compromisso de prazo, orçamento ou capacidade.'),
  table(['Controlo','Valor'],[['Responsável de produto','Sergiu Sandu'],['Versão e data','1.0   '+DATE],['Cadência proposta','Duas semanas por sprint'],['Dimensão','44 histórias e 39 requisitos'],['Entrega local','Validação de fluxos com Sergiu'],['Entrega de produção','Depende de integrações e verificações explícitas']],[140,328]),
  h('Prioridades'),
  p('P0 identifica controlo ou capacidade necessária ao objetivo da fase. P1 identifica melhoria que pode ser adiada se a capacidade for insuficiente. Não atribuir pontos ou dias às histórias sem discussão técnica e validação do âmbito.'),
  h('Responsabilidades da equipa'),
  p('Sergiu prioriza, decide âmbito e aceita comportamento. Desenvolvimento implementa e apresenta evidência. Qualidade valida critérios e regressão relevante. Os representantes de comercial, operações, parque, oficina e financeiro validam o respetivo processo quando forem designados. As posições permanecem por atribuir.'),
  h('Leitura da fase'),
  p('Local corresponde ao âmbito do piloto. Produção exige desenvolvimento e validação adicional. Local e produção contém uma capacidade do piloto e uma extensão posterior explicitada nos critérios. Uma história só é concluída quando a evidência da versão demonstra os critérios; documentação não equivale a implementação.'),
  p('A numeração dos sprints organiza dependências. A indisponibilidade de oficina é uma exceção importante: o bloqueio mínimo de US25 tem de existir antes da aceitação de US14, mesmo que o módulo completo seja aprofundado no sprint 7.')
 ),
 page('Regras de preparação e conclusão',
  h('Antes de iniciar uma história'),
  b('Confirmar quem utiliza a função, que problema resolve e qual é a fase.','Rever requisito, critérios, dados necessários e dependências.','Clarificar regra de autorização e efeito sobre estados, valores e histórico.','Escolher dados de teste sem inventar pessoas reais ou limiares de aprovação.','Dividir história quando integração, segurança ou dados a tornam demasiado grande.'),
  h('Para considerar a história concluída'),
  b('Os critérios de aceitação passam e a evidência está ligada à versão.','As operações são validadas no servidor e respeitam o perfil e o alcance do registo.','Estados, valores e referências mantêm integridade, incluindo repetição da ação.','Interface em português, euros e mensagens compreensíveis.','Persistência e migração são verificadas quando afetadas.','O SOP e o guia do utilizador refletem o comportamento entregue.','Sergiu aceita o resultado da fase e identifica extensões ainda pendentes.'),
  h('Ritmo de cada sprint'),
  p('No planeamento, selecionar histórias pela capacidade e dependências. Durante o trabalho, acompanhar bloqueios e demonstrações curtas. No final, apresentar os critérios com dados de teste, recolher decisão de Sergiu e rever falhas de processo. Transportar histórias incompletas com estado claro; não declarar concluída uma integração apenas desenhada.'),
  h('Evidência mínima'),
  p('Conservar ID da história, versão, cenário, dados utilizados, resultado esperado, resultado observado e responsável da revisão. Para segurança, incluir chamadas diretas à API e testes de acesso alheio. Para financeiro, incluir pagamento parcial, integral, duplicado e superior ao saldo. Para continuidade, incluir recuperação de cópia em ambiente separado.')
 ),
 page('Sequência proposta de sprints',
  table(['Sprint','Objetivo','Histórias'],[[str(i),title,' '.join(s['id'] for s in stories if s['sprint']==i)] for i,title,_,_ in sprints],[49,210,209]),
  h('Dependências externas'),
  b('Escolha de sistema de faturação e validação pelo contabilista antes da integração fiscal real.','Escolha de infraestrutura, base de dados, autenticação e suporte antes do acesso remoto.','Inventário e classificação de dados antes de conservação e eliminação em produção.','Cadastros reais de veículos, motoristas e posições do parque antes da aceitação operacional.','Acesso contratado a telemática, banca ou outros serviços antes de conectores reais.'),
  p('A sequência totaliza aproximadamente vinte semanas se uma equipa tiver capacidade para completar cada sprint. A estimativa deve ser revista depois do primeiro sprint e de conhecer integrações e dados. Segurança, concorrência e integração fiscal podem exigir sprints adicionais.')
 )
]
for sprint,title,goal,exit_ in sprints:
    items=[p('Objetivo  '+goal),p('Revisão de saída  '+exit_)]
    sprint_stories=[x for x in stories if x['sprint']==sprint]
    parts=[sprint_stories[:4],sprint_stories[4:]] if len(sprint_stories)>4 else [sprint_stories]
    for part_index,part in enumerate(parts):
      if part_index: items=[p('Continuação do objetivo  '+goal)]
      for s in part:
        items += [h(s['id']+' '+s['goal']),p('Como '+s['role']+', quero '+s['goal']+' para '+s['benefit']+'.'),p('Módulo '+s['module']+'   Prioridade '+s['priority']+'   Fase '+s['phase']+'   Requisitos '+' '.join(s['requirements']))]
        items += [b(*s['criteria'])]
        if s['dependency']: items += [p('Depende de '+s['dependency'])]
      if sprint==9 and part_index:
          items += [h('Verificações que suportam US37'), b('Entrada e saída, expiração de sessão e desativação de conta.','Motorista tenta ver e alterar serviço alheio; outros perfis tentam decidir sem gerente.','Cotação convertida duas vezes e pagamento com referência repetida não duplicam registos.','Sem prova não há minuta; acima do saldo não há recebimento.','Veículo bloqueado não é despachado nem tem saída operacional.','Autoaprovação normal falha; exceção explícita regista motivo e requerente.','Cópia recuperada mantém relações, histórico e saldos.','Alterações concorrentes e falhas de integração pertencem à aceitação da produção.')]
      PAGES['03_Backlog_e_plano_de_sprints'].append(page('Sprint '+str(sprint)+' '+title+(' continuação' if part_index else ''),*items))
trace=[]
left=requirements[:20];right=requirements[20:]
for i,r in enumerate(left):
    rr=right[i] if i<len(right) else None
    trace.append([r[0],r[4],rr[0] if rr else '',rr[4] if rr else ''])
PAGES['03_Backlog_e_plano_de_sprints'].append(page('Rastreabilidade dos requisitos',
 p('Esta matriz liga o catálogo do documento de requisitos às histórias deste backlog. Cada história identifica também a fase, o perfil e os critérios que permitem validar a entrega.'),
 table(['Requisito','Histórias','Requisito','Histórias'],trace,[68,166,68,166]),
 h('Revisão de alterações'),
 p('Quando um requisito muda, atualizar as histórias ligadas, os procedimentos, o comportamento entregue e os testes afetados. Registar a decisão de Sergiu antes de alterar o âmbito comercial ou as responsabilidades de aprovação.')
))

def style_docx(doc):
    sec=doc.sections[0]; sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.72);sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=sec.footer_distance=Inches(.3)
    for name,size in [('Normal',11),('Title',24),('Heading 1',18),('Heading 2',12)]:
        st=doc.styles[name]; st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_after=Pt(6)
        st.paragraph_format.line_spacing=1.09
        if name!='Normal':st.paragraph_format.keep_with_next=True
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=f.add_run('TSS  |  v1.0  |  ');r.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
    doc.core_properties.title=PAGES[current][0]['title'];doc.core_properties.author='TSS Transportes'
    doc.core_properties.subject='Validação local e preparação da plataforma empresarial'

def docx_table(doc, headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for j,w in enumerate(widths):t.columns[j].width=Inches(w/72)
    for j,head in enumerate(headers):t.rows[0].cells[j].text=head
    pr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');pr.append(repeat)
    for row in rows:
        for j,txt in enumerate(row):t.add_row() if j==0 else None;t.rows[-1].cells[j].text=str(txt)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]/72);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'75');el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcpr.append(borders)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'17354A' if i==0 else ('F2F5F7' if i%2==0 else 'FFFFFF'));tcpr.append(shade)
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(2);par.paragraph_format.line_spacing=1.03
                if j==0 and widths[j]<85:par.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.font.size=Pt(10);run.font.name='Arial';run.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0);run.bold=(i==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

fontbase=Path('C:/Windows/Fonts')
for key,file in [('TSSArial','arial.ttf'),('TSSArialBold','arialbd.ttf')]:pdfmetrics.registerFont(TTFont(key,str(fontbase/file)))
pdfmetrics.registerFontFamily('TSSArial',normal='TSSArial',bold='TSSArialBold',italic='TSSArial',boldItalic='TSSArialBold')
ps={
 'title':ParagraphStyle('title',fontName='TSSArialBold',fontSize=24,leading=28,spaceAfter=15,textColor=colors.black),
 'chapter':ParagraphStyle('chapter',fontName='TSSArialBold',fontSize=18,leading=21,spaceAfter=13,textColor=colors.black),
 'h':ParagraphStyle('h',fontName='TSSArialBold',fontSize=12,leading=15,spaceBefore=9,spaceAfter=5,textColor=colors.black,keepWithNext=True),
 'p':ParagraphStyle('p',fontName='TSSArial',fontSize=11,leading=14.3,spaceAfter=7,textColor=colors.black),
 'b':ParagraphStyle('b',fontName='TSSArial',fontSize=11,leading=14,spaceAfter=4,leftIndent=12,firstLineIndent=-9),
 'cell':ParagraphStyle('cell',fontName='TSSArial',fontSize=10,leading=12.3,spaceAfter=0),
 'head':ParagraphStyle('head',fontName='TSSArialBold',fontSize=10,leading=12.3,textColor=colors.white),
 'short':ParagraphStyle('short',fontName='TSSArial',fontSize=10,leading=12.3,alignment=TA_CENTER)
}
def para(txt,sty='p'):return Paragraph(escape(str(txt)).replace('\n','<br/>'),ps[sty])
def footer(can,doc):
    can.saveState();can.setFont('TSSArial',9);can.setFillColor(colors.HexColor('#444444'));can.drawRightString(540,27,'TSS  |  v1.0  |  '+str(doc.page));can.restoreState()

def build(name,pages):
    doc=Document();style_docx(doc);flow=[]
    for idx,pg in enumerate(pages):
        if idx:doc.add_page_break();flow.append(PageBreak())
        doc.add_paragraph(pg['title'],'Title' if idx==0 else 'Heading 1');flow.append(para(pg['title'],'title' if idx==0 else 'chapter'))
        for item in pg['items']:
            kind=item[0]
            if kind=='p':doc.add_paragraph(item[1]);flow.append(para(item[1]))
            elif kind=='h':doc.add_paragraph(item[1],'Heading 2');flow.append(para(item[1],'h'))
            elif kind=='b':
                for txt in item[1]:
                    dp=doc.add_paragraph('• '+txt);dp.paragraph_format.left_indent=Inches(.15);dp.paragraph_format.first_line_indent=Inches(-.12);dp.paragraph_format.space_after=Pt(3)
                    flow.append(para('• '+txt,'b'))
            elif kind=='table':
                _,headers,rows,widths=item;docx_table(doc,headers,rows,widths)
                data=[[para(x,'head') for x in headers]]+[[para(x,'short' if j==0 and widths[j]<85 else 'cell') for j,x in enumerate(row)] for row in rows]
                tab=Table(data,colWidths=widths,repeatRows=1,hAlign='CENTER')
                tab.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#17354A')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F2F5F7')]),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#D9D9D9')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
                flow += [tab,Spacer(1,8)]
    doc.save(OUT/(name+'.docx'))
    pdf=SimpleDocTemplate(str(OUT/(name+'.pdf')),pagesize=letter,rightMargin=72,leftMargin=72,topMargin=52,bottomMargin=48,title=pages[0]['title'],author='TSS Transportes')
    pdf.build(flow,onFirstPage=footer,onLaterPages=footer)

if __name__=='__main__':
    for current,pages in PAGES.items():
        build(current,pages)
        print(current, 'chapters',len(pages))
    (OUT/'backlog.json').write_text(json.dumps({'version':'1.0','date':'2026-09-30','owner':'Sergiu Sandu','cadence':'Two weeks proposed','stories':stories,'sprints':[dict(number=i,title=t,goal=g,exit=e) for i,t,g,e in sprints]},ensure_ascii=False,indent=2),encoding='utf-8')
    (OUT/'requisitos.json').write_text(json.dumps({'version':'1.0','requirements':[dict(id=a,area=b,description=c,phase=d,stories=e.split()) for a,b,c,d,e in requirements]},ensure_ascii=False,indent=2),encoding='utf-8')
    (OUT/'conteudo_documentos.json').write_text(json.dumps(PAGES,ensure_ascii=False,indent=2),encoding='utf-8')
